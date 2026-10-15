from pptx import Presentation
from pptx.util import Inches
import os

# 设置图片文件夹路径
img_dir = r"E:\xwechat_files\wxid_z9wrrwcf37tu22_71df\msg\file\2025-06\图\图\压缩后"

# 支持的图片格式
img_extensions = ['.jpg', '.jpeg', '.png', '.bmp', '.gif']

# 获取所有图片文件（按文件名排序）
img_files = sorted([
    f for f in os.listdir(img_dir)
    if os.path.splitext(f)[1].lower() in img_extensions
])

# 创建一个新的演示文稿
prs = Presentation()

# 设置幻灯片宽高（默认16:9）
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# 遍历图片并插入到幻灯片中
for img_file in img_files:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局
    img_path = os.path.join(img_dir, img_file)
    slide.shapes.add_picture(img_path, Inches(0), Inches(0),
                             width=prs.slide_width, height=prs.slide_height)

# 保存PPT
output_path = os.path.join(img_dir, "图片汇总.pptx")
prs.save(output_path)

print(f"PPT已保存到：{output_path}")
